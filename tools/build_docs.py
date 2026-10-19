#!/usr/bin/env python
"""把 项目说明文档.md 转成 Word（.docx）。

用法：
    python tools/build_docs.py 项目说明文档.md -o 项目说明文档.docx

支持：# / ## / ### 标题、段落、- 与 1. 列表、> 引用、
``` 代码块、| 表格 |、行内 **粗体** 与 `代码`。
依赖：python-docx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x4F, 0x46, 0xE5)
MUTED = RGBColor(0x64, 0x74, 0x8B)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")


def strip_links(text: str) -> str:
    """把 [文字](链接) 换成 文字（Word 里保留可点击链接意义不大）。"""
    return LINK_RE.sub(r"\1", text)


def set_cn_font(style_or_run, name: str = CN_FONT) -> None:
    """同时设置西文与中文（eastAsia）字体。"""
    font = style_or_run.font
    font.name = name
    run_el = getattr(style_or_run, "element", None)
    if run_el is None:
        return
    rpr = run_el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def shade(paragraph, color: str = "F3F4F6") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


def add_inline(paragraph, text: str, base_bold: bool = False) -> None:
    """把 **粗体** 和 `代码` 转成 run。"""
    text = strip_links(text)
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            set_cn_font(run, CODE_FONT)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)
        else:
            run = paragraph.add_run(chunk)
        if base_bold:
            run.bold = True


def mono_run(paragraph, text: str):
    run = paragraph.add_run(text)
    set_cn_font(run, CODE_FONT)
    run.font.size = Pt(9)
    return run


def mark_header_row(row) -> None:
    """让表头在跨页时重复。"""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def build(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.2)

    # ---- 全局样式 ----
    normal = doc.styles["Normal"]
    set_cn_font(normal)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for level, size in ((1, 18), (2, 14.5), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        set_cn_font(style)
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT if level == 1 else RGBColor(0x0F, 0x17, 0x2A)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(6)

    # ---- 逐行解析 ----
    index = 0
    in_code = False
    code_buffer: list[str] = []
    table_buffer: list[str] = []

    def flush_table() -> None:
        nonlocal table_buffer
        rows = []
        for raw in table_buffer:
            cells = [c.strip() for c in raw.strip().strip("|").split("|")]
            rows.append(cells)
        table_buffer = []
        if not rows:
            return
        # 去掉 |---| 分隔行
        if len(rows) >= 2 and all(set(c) <= set("-: ") for c in rows[1]):
            header, body = rows[0], rows[2:]
        else:
            header, body = None, rows

        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=0, cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        if header:
            header_row = table.add_row()
            mark_header_row(header_row)
            cells = header_row.cells
            for i in range(cols):
                paragraph = cells[i].paragraphs[0]
                add_inline(paragraph, header[i] if i < len(header) else "", True)
                shade(paragraph, "EEEFF5")
        for row in body:
            cells = table.add_row().cells
            for i in range(cols):
                paragraph = cells[i].paragraphs[0]
                add_inline(paragraph, row[i] if i < len(row) else "")
        doc.add_paragraph()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.4)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(8)
                shade(paragraph)
                mono_run(paragraph, "\n".join(code_buffer))
                code_buffer = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_buffer.append(line)
            index += 1
            continue

        # 表格
        if stripped.startswith("|") and stripped.endswith("|"):
            table_buffer.append(stripped)
            index += 1
            continue
        elif table_buffer:
            flush_table()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("---") and set(stripped) <= set("-"):
            index += 1
            continue

        # 标题
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if text:
                paragraph = doc.add_heading(level=min(level, 3))
                add_inline(paragraph, text)
            index += 1
            continue

        # 引用
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.space_after = Pt(3)
            run_text = stripped.lstrip(">").strip()
            add_inline(paragraph, run_text)
            for run in paragraph.runs:
                run.font.color.rgb = MUTED
            index += 1
            continue

        # 无序列表
        if re.match(r"^[-*] ", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:].strip())
            index += 1
            continue

        # 有序列表
        match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if match:
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, match.group(2))
            index += 1
            continue

        # 普通段落
        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    if table_buffer:
        flush_table()
    if code_buffer:
        paragraph = doc.add_paragraph()
        shade(paragraph)
        mono_run(paragraph, "\n".join(code_buffer))

    # ---- 页脚页码 ----
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_end)
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED

    doc.save(str(out_path))
    print(f"[完成] {out_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Markdown 说明书 → Word")
    parser.add_argument("source", help="Markdown 文件")
    parser.add_argument("-o", "--out", default=None, help="输出 docx")
    args = parser.parse_args()

    source = Path(args.source).resolve()
    if not source.is_file():
        print(f"[错误] 文件不存在: {source}")
        return 2
    out = Path(args.out).resolve() if args.out else source.with_suffix(".docx")
    build(source, out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
